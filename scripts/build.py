"""Build the website, source library and matching Word/PDF plan from one manuscript."""
from pathlib import Path
from html import escape
from urllib.parse import quote,urlsplit,unquote
import json,re,hashlib,zipfile,argparse,subprocess,sys
ROOT=Path(__file__).resolve().parents[1]
BASE='https://auraofintelligence.github.io/installing-the-missing-middle/'
REPO='https://github.com/auraofintelligence/installing-the-missing-middle'
SLUGS=['index','federation','national-discovery','installation','learning','workloads','finance','economics','renewal','law-and-stewardship','component-briefs','next-development']
SHORT=['The invitation','The federation','National discovery','Installation & heat','Workforce learning','What runs here','Finance & ownership','The numbers','Generational renewal','Law & stewardship','Component briefs','Next development']
CSS=(ROOT/'assets/site.css').read_text()
JS=(ROOT/'assets/site.js').read_text()
def load_plan():
 raw=(ROOT/'content/plan.md').read_text();chapters=[]
 for idx,block in enumerate(re.split(r'^## ',raw,flags=re.M)[1:]):
  first,*rest=block.split('\n',1);body=rest[0];chunks=re.split(r'^### ',body,flags=re.M);lead=chunks[0].strip();parts=[];references=[]
  for ch in chunks[1:]:
   title,txt=ch.split('\n',1)
   if '\nSources: ' in txt:txt,rr=txt.rsplit('\nSources: ',1);references=re.findall(r'[DP]\d\d',rr)
   parts.append({'title':title.strip(),'text':txt.strip()})
  chapters.append({'slug':SLUGS[idx],'title':first.strip(),'lead':lead,'parts':parts,'refs':references})
 return chapters,raw
CH,RAW=load_plan();S=json.loads((ROOT/'content/sources.json').read_text());DOCS=S['documents'];LINKS=S['links'];REF={x['id']:x for x in DOCS+LINKS};ART={x['slug']:x for x in json.loads((ROOT/'content/artwork.json').read_text()) if x.get('path')} if (ROOT/'content/artwork.json').exists() else {}
NAV=[(c['slug']+'.html',SHORT[i]) for i,c in enumerate(CH)]+[('sources.html','Document library'),('connections.html','Project connections'),('downloads.html','Downloads'),('licence.html','Licence'),('sitemap.html','All pages')]
def anchor(t):return re.sub(r'[^a-z0-9]+','-',t.lower()).strip('-')
def fmt(t):
 s=escape(t)
 return re.sub(r'\b([DP]\d\d)\b',lambda m:f'<a href="sources.html#{m[1]}" title="{escape(REF.get(m[1],{}).get("title",m[1]),quote=True)}">{m[1]}</a>',s)
def refs(ids):return '<div class="source-chips">'+''.join(f'<a class="chip" href="sources.html#{id}">{id} · {escape(REF[id]["title"])}</a>' for id in ids if id in REF)+'</div>'
def hero(slug):
 a=ART.get(slug)
 if not a:return ''
 from PIL import Image
 w,h=Image.open(ROOT/a['path']).size
 return f'<figure class="hero"><img src="{a["path"]}" alt="{escape(a.get("alt","Concept artwork"),quote=True)}" width="{w}" height="{h}" fetchpriority="high"><figcaption>Concept artwork · <a href="{a["source_page"]}" target="_blank" rel="noopener noreferrer">Connected source project</a>. Illustrative, not a completed installation.</figcaption></figure>'
def shell(slug,title,lead,body,art=True,section='The plan',description=None):
 filename=slug+'.html';names=[x[0] for x in NAV];pos=names.index(filename) if filename in names else 0;prev=NAV[(pos-1)%len(NAV)];nxt=NAV[(pos+1)%len(NAV)]
 nav=''.join(f'<a href="{u}"'+(' aria-current="page"' if u==filename else '')+f'>{escape(n)}</a>' for u,n in NAV)
 actions='<div class="actions"><a class="button" href="presentation.html">Open the presentation</a><a class="button secondary" href="downloads/Installing-the-Missing-Middle-Plan.pdf">Read the PDF</a><a class="button secondary" href="downloads/Installing-the-Missing-Middle-Plan.docx">Editable Word plan</a></div>' if slug=='index' else ''
 return f'''<!doctype html><html lang="en-AU"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(title)} | Installing the Missing Middle</title><meta name="description" content="{escape(description or lead,quote=True)}"><meta name="theme-color" content="#203b4c"><link rel="canonical" href="{BASE+('' if slug=='index' else filename)}"><link rel="icon" href="assets/favicon.png" type="image/png"><link rel="stylesheet" href="assets/site.css"><script defer src="assets/site.js"></script><link rel="describedby" href="llms.txt"></head><body id="top"><a class="skip" href="#main">Skip to content</a><header class="header"><div class="wrap header-inner"><a class="brand" href="index.html"><span class="brand-mark" aria-hidden="true">M</span><span>Installing the<br>Missing Middle</span></a><nav class="main-nav" aria-label="Main navigation"><a href="index.html">The plan</a><a href="component-briefs.html">Work together</a><a href="sources.html">Sources</a><button class="menu-button" data-menu aria-expanded="false" aria-controls="all-nav">All chapters</button></nav></div><nav class="all-nav" id="all-nav" aria-label="All chapters" hidden><div class="wrap nav-grid">{nav}<a href="presentation.html">Presentation view</a></div></nav></header><main id="main"><div class="wrap intro"><p class="eyebrow">{escape(section)} · Luke Nathan Hayes</p><h1>{escape(title)}</h1><p class="lead">{escape(lead)}</p>{actions}</div>{hero(slug) if art else ''}{body}</main><footer class="footer"><div class="wrap route-nav"><a href="{prev[0]}">Previous<strong>← {escape(prev[1])}</strong></a><a href="{nxt[0]}">Next<strong>{escape(nxt[1])} →</strong></a></div><div class="wrap footer-base"><p><strong>A federation of people and corporate vehicles.</strong><br>© 2026 Luke Nathan Hayes · Strange But True Public Source Licence.<br>Working edition 1.0 · 29 October 2026</p><div class="footer-links"><a href="sitemap.html">All pages</a><a href="downloads.html">PDF & Word</a><a href="licence.html">Licence</a><a href="{REPO}" target="_blank" rel="noopener noreferrer">GitHub ↗</a></div></div></footer><button class="top" data-top aria-label="Return to top">↑</button></body></html>'''
def write(path,text):
 p=ROOT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding='utf-8')
def build_web():
 from PIL import Image,ImageDraw,ImageFont
 im=Image.new('RGB',(128,128),'#203b4c');d=ImageDraw.Draw(im);f=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',78) if Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf').exists() else ImageFont.load_default();d.text((64,62),'M',font=f,anchor='mm',fill='white');im.save(ROOT/'assets/favicon.png')
 for i,c in enumerate(CH):
  article=''.join(f'<section id="{anchor(p["title"])}"><h2>{escape(p["title"])}</h2><p>{fmt(p["text"])}</p></section>' for p in c['parts'])
  aside='<aside class="aside"><h2>On this page</h2>'+''.join(f'<a href="#{anchor(p["title"])}">{escape(p["title"])}</a>' for p in c['parts'])+'</aside>'
  metric='<div class="metrics"><div class="metric"><strong>30 million</strong><span>Future population planning basis</span></div><div class="metric"><strong>3,000 nodes</strong><span>One rack per 10,000 people</span></div><div class="metric"><strong>Every postcode</strong><span>In the invitation from the outset</span></div></div>' if i==0 else ''
  extra=''
  if c['slug']=='economics':extra='''<div class="calc" data-calc><h2>Explore the initial capital scenario</h2><p>A population-and-cost illustration, not a deployment schedule or investment forecast. Inputs remain in this page only.</p><div class="calc-grid"><label for="population">Planning population<input id="population" type="number" min="1" step="1000" value="30000000"></label><label for="ratio">People per node<input id="ratio" type="number" min="1" step="100" value="10000"></label><label for="cost">All-in site allowance (A$)<input id="cost" type="number" min="0" step="100000" value="6000000"></label></div><output id="calc-result" aria-live="polite"></output><p>Nodes round upwards. Initial capital excludes subsequent operations, financing and refresh cycles.</p></div>'''
  if c['slug']=='component-briefs':extra='<div class="actions">'+''.join(f'<a class="button secondary" href="briefs/{x.lower()}.md" download>{x}’s agent brief</a>' for x in ['Brad','Megan','Luke'])+'</div><p><a href="AGENTS.md">Shared agent instructions</a> · <a href="briefs/current-direction.md">Current direction</a></p>'
  body=f'<div class="wrap">{metric}<div class="content-layout"><article class="reading">{article}{extra}{refs(c["refs"]) if i else ""}</article>{aside}</div></div>'
  write(c['slug']+'.html',shell(c['slug'],c['title'],c['lead'],body))
 cards=[]
 for d in DOCS:
  exists=(ROOT/d['path']).exists();name=quote(d['path'],safe='/')
  actions=(f'<a class="button secondary" href="{name}">Original {Path(d["path"]).suffix[1:].upper()}</a> <a href="documents/readers/{d["id"]}.html">Read online</a>') if exists else '<p>Original supplied and indexed. This exact edition is awaiting repository transfer; it has not been replaced by a different version.</p>'
  cards.append(f'<article class="card" id="{d["id"]}" data-search="{escape(d["title"]+" "+d["original_name"],quote=True)}"><p class="tag">{d["id"]} · Original document</p><h2>{escape(d["title"])}</h2><p>{escape(d.get("note","Source document retained for its project context and detail."))}</p><p class="metadata">{escape(d["original_name"])}</p>{actions}<details><summary>Original-file verification</summary><p class="hash">SHA-256<br>{d.get("sha256") or "Pending exact-file transfer"}</p></details></article>')
 project_cards=[]
 for l in LINKS:project_cards.append(f'<article class="card" id="{l["id"]}" data-search="{escape(l["title"]+" "+l["purpose"],quote=True)}"><p class="tag">{l["id"]} · Supplied project</p><h2><a href="{l["url"]}" target="_blank" rel="noopener noreferrer">{escape(l["title"])} ↗</a></h2><p>{escape(l["purpose"])}</p></article>')
 controls='<div class="filter"><label class="sr-only" for="search">Search the library</label><input id="search" data-filter type="search" placeholder="Find a document or project"></div><p data-count aria-live="polite">28 entries</p><p data-empty class="empty" hidden>No matches. Try a broader term.</p>'
 intro='<p>The register covers all 14 supplied documents. Available original files are preserved unchanged. The project links supplied during development are collected below. The current plan and direction register govern new work; older originals retain their original wording.</p>'
 body='<section class="wrap wide-section">'+intro+controls+'<div class="card-grid">'+''.join(cards+project_cards)+'</div><div class="notice"><p>Original sources contain their own external references and rights. <a href="content/embedded-links.json">All extracted document links</a> are also available as a machine-readable register. This register preserves references; it is not a new verification of every linked claim.</p></div></section>'
 write('sources.html',shell('sources','The complete source library','Original documents, project connections and a clear path back to the material.',body,section='Read the sources'))
 body='<section class="wrap wide-section"><p>Every project supplied directly in this conversation, with its role in the proposal. Links connect distinct projects; they do not merge their organisations or permissions.</p><div class="card-grid">'+''.join(project_cards)+'</div></section>'
 write('connections.html',shell('connections','The work already connected','Thirteen supporting project connections and the presentation repository.',body,section='Project connections'))
 body='''<section class="wrap wide-section"><div class="card-grid"><article class="card"><h2>The business plan</h2><p>Matching editions built from the same manuscript.</p><div class="actions"><a class="button" href="downloads/Installing-the-Missing-Middle-Plan.pdf">PDF plan</a><a class="button secondary" href="downloads/Installing-the-Missing-Middle-Plan.docx">Editable Word</a><a href="content/plan.md">Markdown manuscript</a></div></article><article class="card"><h2>The source pack</h2><p>The original uploads available in the repository, preserved without editing, plus the source and project-link registers. The library identifies the status of each exact edition.</p><div class="actions"><a class="button secondary" href="downloads/Installing-the-Missing-Middle-Source-Pack.zip">Download source pack</a><a href="sources.html">Browse online</a></div></article><article class="card"><h2>Component handovers</h2><p>Separate starting briefs for the three complementary streams.</p><p><a href="briefs/brad.md">Brad</a> · <a href="briefs/megan.md">Megan</a> · <a href="briefs/luke.md">Luke</a></p><p><a href="AGENTS.md">Agent instructions</a> · <a href="briefs/current-direction.md">Current direction</a></p></article><article class="card"><h2>Present and continue</h2><p><a href="presentation.html">Open presentation view</a> for a chapter-by-chapter conversation. Read the full chapters for detail.</p><p><a href="'''+REPO+'''">Source repository</a> contains the content and build process. No accounts, trackers or submission backend are added to this site.</p></article></div></section>'''
 write('downloads.html',shell('downloads','Read, share and develop the plan','PDF for reading. Word for editing. Markdown and original sources for the agents.',body,section='Downloads'))
 lic=(ROOT/'LICENCE.md').read_text();licbody=''.join('<h2>'+escape(x[3:])+'</h2>' if x.startswith('## ') else '<p>'+escape(x)+'</p>' for x in lic.split('\n\n') if x and not x.startswith('# '))
 write('licence.html',shell('licence','Strange But True Public Source Licence','The established public-source terms. Commercial and corporate rights remain with Luke Nathan Hayes.','<div class="wrap content-layout"><article class="reading">'+licbody+'<p><a href="LICENCE.md">Read the licence file</a> · <a href="RIGHTS.md">Source and artwork rights</a></p></article></div>',section='Rights and attribution'))
 toc='<ol class="toc-list">'+''.join(f'<li><a href="{u}"><strong>{escape(t)}</strong>{escape(CH[i]["lead"]) if i<len(CH) else ""}</a></li>' for i,(u,t) in enumerate(NAV))+'<li><a href="presentation.html">Presentation view</a></li></ol>'
 write('sitemap.html',shell('sitemap','Choose your next chapter','The complete plan, its sources and the material for continuing the work.','<section class="wrap wide-section">'+toc+'</section>',section='All pages'))
 slides=''.join(f'<section class="presentation-card" id="slide-{i+1}"><p>{i+1:02} / {len(CH)}</p><h2>{escape(c["title"])}</h2><p>{escape(c["lead"])}</p><p>{escape(c["parts"][0]["text"])}</p><a href="{c["slug"]}.html">Read this chapter →</a></section>' for i,c in enumerate(CH))
 controls='<div class="presentation-controls"><button data-present>Toggle presentation mode</button><button data-prev aria-label="Previous slide">← Previous</button><span data-slide-count aria-live="polite"></span><button data-next aria-label="Next slide">Next →</button></div><p>Use the arrow keys in presentation mode. Escape returns to the full outline.</p>'
 write('presentation.html',shell('presentation','Installing the Missing Middle','A conversation about the business, the federation and the components we build together.','<section class="wrap wide-section">'+controls+slides+'</section>',art=False,section='Presentation'))
 write('404.html',shell('404','Let’s find the right chapter','The page may have moved. The plan and complete source library are here.','<section class="wrap wide-section"><a class="button" href="index.html">Open the plan</a> <a href="sitemap.html">All pages</a></section>',art=False))
 write('.nojekyll','');write('robots.txt','# Project-path file only. Origin-wide crawler rules belong at https://auraofintelligence.github.io/robots.txt\n# Public discovery does not grant reuse or training rights beyond LICENCE.md.\nSitemap: '+BASE+'sitemap.xml\n')
 write('llms.txt','# Installing the Missing Middle\n\nBusiness plan by Luke Nathan Hayes, working edition 1.0, 29 October 2026.\n\nRead content/plan.md and briefs/current-direction.md before interpreting the archive. Preserve the federation, finite investor returns, co-operative ownership, personal data ownership and the roles described. All original and third-party rights remain in force. Discovery is not a blanket AI-training permission.\n\n'+'\n'.join(f'- [{t}]({BASE+u})' for u,t in NAV)+'\n\n- [Manuscript]('+BASE+'content/plan.md)\n- [Licence]('+BASE+'LICENCE.md)\n')
 urls=[BASE+('' if u=='index.html' else u) for u,_ in NAV]+[BASE+'presentation.html']
 write('sitemap.xml','<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'+''.join('<url><loc>'+escape(u)+'</loc></url>' for u in urls)+'</urlset>')
def readers():
 import fitz
 from docx import Document
 from docx.text.paragraph import Paragraph
 from docx.table import Table
 extracted=[]
 for d in DOCS:
  p=ROOT/d['path']
  if not p.exists():continue
  blocks=[]
  if p.suffix.lower()=='.pdf':
   doc=fitz.open(p)
   for i,page in enumerate(doc):
    blocks.append(f'<h2 class="page-label">Page {i+1} of {len(doc)}</h2><pre>{escape(page.get_text())}</pre>')
    for lk in page.get_links():
     if lk.get('uri','').startswith(('https://','http://')):extracted.append({'document':d['id'],'page':i+1,'url':lk['uri']})
  else:
   doc=Document(p)
   for child in doc.element.body:
    if child.tag.endswith('}p'):
     q=Paragraph(child,doc);text=q.text
     if text:blocks.append('<p>'+escape(text)+'</p>')
    elif child.tag.endswith('}tbl'):
     tb=Table(child,doc);blocks.append('<div class="table-scroll"><table>'+''.join('<tr>'+''.join('<td>'+escape(cell.text)+'</td>' for cell in row.cells)+'</tr>' for row in tb.rows)+'</table></div>')
   for rel in doc.part.rels.values():
    if rel.is_external and str(rel.target_ref).startswith(('http://','https://')):extracted.append({'document':d['id'],'url':rel.target_ref})
  html='<div class="wrap"><article class="reader"><p><a href="'+quote(d['path'],safe='/')+'">Open the unchanged original</a></p><div class="notice">This is a searchable reading copy. The original retains its page layout, figures, links and authoritative wording.</div>'+''.join(blocks)+'</article></div>'
  full=shell(d['id'],d['title'],'Original source reading copy.',html,art=False,section='Document library')
  full=full.replace('<head>','<head><base href="../../">',1).replace(BASE+d['id']+'.html',BASE+'documents/readers/'+d['id']+'.html')
  write('documents/readers/'+d['id']+'.html',full)
 write('content/embedded-links.json',json.dumps(extracted,ensure_ascii=False,indent=2)+'\n')
def handovers():
 shared='''# Shared agent instructions

Develop the named human lead’s component of Installing the Missing Middle. Read content/plan.md, briefs/current-direction.md and the relevant source documents first.

Use Australian English, Australian institutional context, AUD and metric units. Verify any operative legal, funding, procurement or technical claim against the applicable current primary source when the component needs it; do not import US defaults.

Preserve the federation of people and corporate vehicles. Do not impose a compulsory parent company, permanent extraction model, founder veto or a requirement that all participants become co-operatives. Keep co-operative ownership as the asset pathway specified in the plan.

Brad builds the core infrastructure delivery team through his industry knowledge and contacts. Megan develops institutional and commercial arrangements. Luke carries the integrated architecture, applications and progression plans. Do not reduce these contributions to incidental introductions or repeat a lecture about the present stage of Luke’s systems.

Current human directions govern new work where an archive differs. C-Hour is non-monetary. Personal Auras and data belong to their people. Keep voluntary contribution, paid work and investment in distinct records.

Return a usable component, sources, explicit assumptions, dependencies on the other components and clearly identified proposals for decisions. Specific evidenced technical or commercial challenges are welcome; silently changing the purpose is not.

External outreach, expenditure, agreements, public statements or changes to other repositories require authority from the responsible person. This brief alone does not authorise autonomous external actions.
'''
 write('AGENTS.md',shared)
 direction='''# Current direction — 29 October 2026

This record preserves Luke’s directions in this conversation for the plan and subsequent agent work.

- The business focuses on installing single-rack frontier compute, UPS and useful heat recovery; maintaining, upgrading and redeploying equipment remain part of the business.
- Future population basis: 30 million; one rack per 10,000 people; 3,000-node planning baseline. Every postcode enters outreach in parallel. Actual installations follow the participating people, host facilities, workloads and agreements.
- ready SET Co-op seeks members; P4A develops council/division and state/federal electorate twin connections. Identify schools, hospitals, pools, wellbeing and aged-care/longevity hubs in those regions.
- Brad coordinates data-centre UPS installers, knows the industry from single racks to the latest public/private facilities, and has the contacts to build the core team. Megan brings hospital contract management, contract-law understanding and a former HIMAA role.
- The structure is a federation of people and corporate vehicles, not one mandatory corporate parent.
- Finance combines public support, finite-return investment and co-operative acquisitions. Acquisition payments can replenish deployment capital. Investor return and installer margin are separate.
- Capture installations in 360 VR and wearable XR, build environment models and just-in-time learning, and retain the contribution of skilled trades and mentors.
- Luke has progression plans for the legal engine, Aura and connected work. Expand them as required under his direction; do not substitute an agent’s assumptions for the architecture.
- The supplied programmes describe part of the workload base, alongside electorate simulations and whatever communities choose to develop with personal Auras and connected AI tools.
- Additional public-purpose applications create funding workstreams. No award, institution’s endorsement or family member’s formal appointment is invented.
- C-Hour remains non-monetary; no conversion into dollars or ownership.
- Preserve original documents and their rights unchanged. The website uses clear navigation, large uncropped heroes, previous/next routes, a floating top button, source links, PDF/Word downloads and the established Strange But True Public Source Licence.
'''
 write('briefs/current-direction.md',direction)
 for person,idx in [('Brad',0),('Megan',1),('Luke',2)]:
  part=CH[10]['parts'][idx]
  txt=f'# {person} — component development brief\n\nInstalling the Missing Middle | Luke Nathan Hayes | 29 October 2026\n\n## The component\n\n{part["text"]}\n\n## Shared interfaces\n\n{CH[10]["parts"][3]["text"]}\n\n## Read first\n\n- ../content/plan.md\n- current-direction.md\n- ../AGENTS.md\n- ../content/sources.json\n\n## Return\n\nA developed component, its cost and timing assumptions where relevant, source-backed findings, dependencies, and proposals for the people responsible to decide. Distinguish confirmed inputs from choices still being developed.\n'
  write('briefs/'+person.lower()+'.md',txt)
 write('briefs/site-brief.md','# Candidate installation brief\n\nUse with the plan; this is a development template, not a submitted application.\n\n'+'\n\n'.join('## '+x+'\n\nTo develop with the relevant participants.' for x in ['Place and geographic memberships','Host, owner and decision-makers','Participating entities and co-operative pathway','Workloads, users and prospective customers','Rack generation and supporting systems','Heat use and building interfaces','Power, UPS and continuity','Connectivity, storage and permissions','Physical works and delivery team','Commercial terms, funding and cost allocation','Learning capture and workforce','Acceptance, handover and maintenance','Acquisition, investor exit and renewal','Sources, assumptions and next actions']))
def docx_build():
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.enum.text import WD_ALIGN_PARAGRAPH
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.05;normal.paragraph_format.space_after=Pt(5)
 for sn,size,col in [('Title',34,'203B4C'),('Subtitle',15,'42676B'),('Heading 1',22,'203B4C'),('Heading 2',13,'285B62')]:
  s=doc.styles[sn];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(col);s.paragraph_format.space_before=Pt(7);s.paragraph_format.space_after=Pt(4)
 for sn in ['Heading 1','Heading 2']:doc.styles[sn].paragraph_format.keep_with_next=True
 foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('Installing the Missing Middle | Luke Nathan Hayes | ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 for r in foot.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor.from_string('526568')
 header=sec.header.paragraphs[0];header.add_run('WORKING DEVELOPMENT EDITION 1.0  /  29 OCTOBER 2026').font.size=Pt(8.5)
 def link(par,text,url):
  rId=par.part.relate_to(url,RT.HYPERLINK,is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rId);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'285B62');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);el.append(r);par._p.append(el)
 doc.add_paragraph('INSTALLING THE\nMISSING MIDDLE','Title');doc.add_paragraph('Community compute. Useful heat. Local ownership.','Subtitle');doc.add_paragraph('A business plan for a federation of people and corporate vehicles.')
 a=ART.get('index')
 if a and (ROOT/a['path']).exists():
  from PIL import Image
  import io
  picture=io.BytesIO();Image.open(ROOT/a['path']).convert('RGB').save(picture,format='JPEG',quality=92);picture.seek(0)
  doc.add_picture(picture,width=Inches(6.65));q=doc.add_paragraph('Concept artwork from the ready SET Co-op Cultural Intelligence Node.');q.runs[0].font.size=Pt(9)
 doc.add_paragraph('Prepared by Luke Nathan Hayes\nFor joint development with Megan and Brad\n29 October 2026')
 doc.add_paragraph('The national starting point','Heading 2');doc.add_paragraph('30 million people · One single-rack node per 10,000 people · 3,000-node planning baseline')
 doc.add_paragraph('The installation business connects practical delivery, useful heat, workforce development and ongoing upgrades with a financed pathway into local co-operative ownership.')
 q=doc.add_paragraph();link(q,'Presentation and complete source library',BASE);doc.add_paragraph('Original material: Strange But True Public Source Licence. Third-party material retains its own rights.')
 doc.add_heading('Inside the plan',level=1).paragraph_format.page_break_before=True
 for i,c in enumerate(CH):doc.add_paragraph(f'{i+1:02}  {c["title"]}')
 doc.add_paragraph('The complete source library follows the plan: all 14 original-upload records, all supplied project links, component handovers and the established licence.')
 for i,c in enumerate(CH):
  doc.add_heading(f'{i+1:02}  {c["title"]}',level=1).paragraph_format.page_break_before=True;doc.add_paragraph(c['lead'],'Subtitle')
  for s in c['parts']:doc.add_heading(s['title'],level=2);doc.add_paragraph(s['text'])
  q=doc.add_paragraph('Sources: ')
  for j,id in enumerate(c['refs']):
   if j:q.add_run(' · ')
   link(q,id,BASE+'sources.html#'+id)
 for start in [0,7]:
  doc.add_heading('Original document library'+(' — continued' if start else ''),level=1).paragraph_format.page_break_before=True
  if start==0:doc.add_paragraph('Original files are preserved unchanged. Current directions govern the new plan where archive wording differs. The complete website library identifies the status of every supplied edition and provides searchable reading copies and original downloads.')
  for d in DOCS[start:start+7]:
   doc.add_heading(d['id']+'  '+d['title'],level=2);doc.add_paragraph(d.get('note','Source document supplied for this proposal.'));q=doc.add_paragraph();link(q,'Original file and reading copy',BASE+'sources.html#'+d['id'])
 for start in [0,7]:
  doc.add_heading('Supplied project connections'+(' — continued' if start else ''),level=1).paragraph_format.page_break_before=True
  for l in LINKS[start:start+7]:
   doc.add_heading(l['id']+'  '+l['title'],level=2);doc.add_paragraph(l['purpose']);q=doc.add_paragraph();link(q,l['url'],l['url'])
 doc.add_heading('Continue the work with your agents',level=1).paragraph_format.page_break_before=True
 doc.add_paragraph('Separate briefs develop the three components while retaining the shared direction. The repository contains the full manuscript, source register, original uploads and build process.')
 for name in ['brad','megan','luke']:
  q=doc.add_paragraph();link(q,name.title()+' — component brief',BASE+'briefs/'+name+'.md')
 for name,url in [('Shared agent instructions','AGENTS.md'),('Current direction','briefs/current-direction.md'),('Candidate site brief','briefs/site-brief.md')]:
  q=doc.add_paragraph();link(q,name,BASE+url)
 doc.add_heading('Publication and source notes',level=2);doc.add_paragraph('The Word edition, PDF and website are generated from content/plan.md. The original source documents retain their original terms, claims, dates and development history. The 29 October 2026 direction supplies the federation structure, family contributions, parallel national outreach and finite-return ownership model for this plan.')
 doc.add_paragraph('Artwork is illustrative material from connected source projects, with provenance retained in content/artwork.json. No image establishes a completed installation, institutional endorsement or a supplier commitment.')
 for block in (ROOT/'LICENCE.md').read_text().split('\n\n'):
  if block.startswith('# '):doc.add_heading(block[2:],level=1).paragraph_format.page_break_before=True
  elif block.startswith('## '):doc.add_heading(block[3:],level=2)
  elif block:doc.add_paragraph(block)
 doc.core_properties.title='Installing the Missing Middle';doc.core_properties.subject='Business plan for community compute, useful heat and co-operative ownership';doc.core_properties.author='Luke Nathan Hayes';doc.core_properties.keywords='community compute; useful heat; ready SET Co-op; federation; Megan; Brad'
 path=ROOT/'downloads/Installing-the-Missing-Middle-Plan.docx';path.parent.mkdir(exist_ok=True);doc.save(path);return path

def checks():
 from html.parser import HTMLParser
 class Links(HTMLParser):
  def __init__(self):super().__init__();self.links=[];self.ids=set();self.base=False
  def handle_starttag(self,t,a):
   d=dict(a)
   if 'id' in d:self.ids.add(d['id'])
   if t=='base':self.base=True
   for k in ['href','src']:
    if k in d:self.links.append(d[k])
 errors=[];count=0
 for f in ROOT.rglob('*.html'):
  if '.git' in f.parts:continue
  parser=Links();parser.feed(f.read_text());count+=1
  for link in parser.links:
   u=urlsplit(link)
   if u.scheme or u.netloc or not u.path:continue
   target=(ROOT if parser.base else f.parent)/unquote(u.path)
   if not target.exists():errors.append(str(f.relative_to(ROOT))+': '+link)
 for d in DOCS:
  p=ROOT/d['path']
  if p.exists():
   b=p.read_bytes();sha=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
   if sha!=d['git_blob']:errors.append('Original changed: '+d['id'])
 report={'html_pages':count,'originals_present':sum((ROOT/d['path']).exists() for d in DOCS),'originals_expected':14,'supplied_project_links':len(LINKS),'chapters':len(CH),'errors':errors,'plan_sha256':hashlib.sha256((ROOT/'content/plan.md').read_bytes()).hexdigest()}
 write('content/build-report.json',json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2));return not errors

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--pdf',action='store_true');ap.add_argument('--check',action='store_true');args=ap.parse_args()
 handovers();readers();build_web();path=docx_build()
 if args.pdf:subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/missing-middle-lo','--headless','--convert-to','pdf','--outdir',str(path.parent),str(path)],check=True,timeout=120)
 with zipfile.ZipFile(ROOT/'downloads/Installing-the-Missing-Middle-Source-Pack.zip','w',zipfile.ZIP_DEFLATED) as z:
  for d in DOCS:
   p=ROOT/d['path']
   if p.exists():z.write(p,'originals/'+p.name)
  for f in ['content/sources.json','content/embedded-links.json','content/artwork.json','LICENCE.md','RIGHTS.md']:
   if (ROOT/f).exists():z.write(ROOT/f,f)
 write('downloads/plan-version.json',json.dumps({'edition':'1.0','date':'2026-10-29','source_sha256':hashlib.sha256((ROOT/'content/plan.md').read_bytes()).hexdigest(),'word':'Installing-the-Missing-Middle-Plan.docx','pdf':'Installing-the-Missing-Middle-Plan.pdf'},indent=2)+'\n')
 if args.check and not checks():sys.exit(1)
if __name__=='__main__':main()
